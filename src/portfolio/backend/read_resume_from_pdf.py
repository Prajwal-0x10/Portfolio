import os
from pathlib import Path
from typing import List, Dict, Union
from pypdf import PdfReader
from docx import Document


def read_pdf_text(file_path: Union[str, Path]) -> str:
    """
    Reads and extracts text content from a PDF file given its file path.

    Args:
        file_path (str | Path): Path to the target PDF file.

    Returns:
        str: Cleaned text concatenated from all pages of the PDF.

    Raises:
        FileNotFoundError: If the file does not exist at the given path.
        ValueError: If the file is not a valid PDF file.
    """
    path = Path(file_path)

    if not path.exists():
        raise FileNotFoundError(f"File not found at path: {path.resolve()}")

    if path.suffix.lower() != ".pdf":
        raise ValueError(f"File at {path} is not a PDF file (extension: '{path.suffix}')")

    reader = PdfReader(path)
    text_list = []

    for i, page in enumerate(reader.pages):
        page_text = page.extract_text()
        if page_text:
            text_list.append(page_text.strip())

    return "\n\n".join(text_list)


def read_docx_text(file_path: Union[str, Path]) -> str:
    """
    Reads and extracts text content from a DOCX file given its file path.

    Args:
        file_path (str | Path): Path to the target DOCX file.

    Returns:
        str: Cleaned text extracted from paragraphs and tables.

    Raises:
        FileNotFoundError: If the file does not exist at the given path.
        ValueError: If the file is not a valid DOCX file.
    """
    path = Path(file_path)

    if not path.exists():
        raise FileNotFoundError(f"File not found at path: {path.resolve()}")

    if path.suffix.lower() != ".docx":
        raise ValueError(f"File at {path} is not a DOCX file (extension: '{path.suffix}')")

    doc = Document(path)
    text_list = []

    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        if text:
            text_list.append(text)

    for table in doc.tables:
        for row in table.rows:
            row_text = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if row_text:
                text_list.append(" | ".join(row_text))

    return "\n\n".join(text_list)


def read_resume_text(file_path: Union[str, Path]) -> str:
    """
    Reads text content from a resume file automatically based on its extension (.pdf or .docx).

    Args:
        file_path (str | Path): Path to the PDF or DOCX file.

    Returns:
        str: Extracted text content.

    Raises:
        FileNotFoundError: If the file does not exist.
        ValueError: If the file format is not supported (.pdf or .docx).
    """
    path = Path(file_path)
    ext = path.suffix.lower()

    if ext == ".pdf":
        return read_pdf_text(path)
    elif ext == ".docx":
        return read_docx_text(path)
    else:
        raise ValueError(f"Unsupported file format '{ext}'. Only .pdf and .docx files are supported.")


if __name__ == "__main__":
    import sys

    if sys.platform == "win32":
        sys.stdout.reconfigure(encoding='utf-8')

    if len(sys.argv) > 1:
        file_path = sys.argv[1]
        try:
            extracted_text = read_resume_text(file_path)
            print(f"--- Extracted Text from '{file_path}' ---")
            print(extracted_text)
        except Exception as e:
            print(f"Error reading file: {e}")
    else:
        print("Usage: python read_resume_from_pdf.py <path_to_pdf_or_docx>")
